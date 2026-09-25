# -*- coding: utf-8 -*-
"""文档预览：docx→HTML / Word COM→PDF。表格由前端 SheetJS 预览。"""

from __future__ import annotations

import hashlib
import os

from .paths import IS_WINDOWS

try:
    from docx import Document as DocxDocument

    HAS_DOCX = True
except ImportError:
    HAS_DOCX = False

try:
    import mammoth

    HAS_MAMMOTH = True
except ImportError:
    HAS_MAMMOTH = False

try:
    import win32com.client

    HAS_WIN32 = True
except ImportError:
    HAS_WIN32 = False

_pdf_cache_dir: str | None = None
_word_available: bool | None = None


def extract_docx_html(filepath: str) -> tuple[str, bool]:
    """
    将 docx 转为带样式的 HTML。
    优先 mammoth，回退 python-docx 纯文本。
    返回 (html_or_text, is_html)。
    """
    if HAS_MAMMOTH:
        try:
            with open(filepath, "rb") as f:
                result = mammoth.convert_to_html(f)
                html_out = result.value
                if html_out.strip():
                    return html_out, True
                return "（文档为空或无可显示的内容）", False
        except Exception:
            pass

    if not HAS_DOCX:
        return "（python-docx 和 mammoth 库均未安装，无法预览）", False

    try:
        doc = DocxDocument(filepath)
        lines = []
        for para in doc.paragraphs:
            text = para.text.strip()
            if text:
                if para.style and para.style.name:
                    sn = para.style.name
                    if "Heading" in sn or "Title" in sn:
                        lines.append(f"【{sn}】{text}")
                    else:
                        lines.append(text)
                else:
                    lines.append(text)
        for i, table in enumerate(doc.tables):
            lines.append(f"\n--- 表格 {i + 1} ---")
            for row in table.rows:
                cells = [cell.text.strip() for cell in row.cells]
                lines.append(" | ".join(cells))
            lines.append("")
        if not lines:
            return "（文档为空或无可显示的文本内容）", False
        return "\n".join(lines), False
    except Exception as e:
        return f"读取文件失败: {e}", False


def get_pdf_cache_dir() -> str:
    global _pdf_cache_dir
    if _pdf_cache_dir is None:
        cache_base = os.path.join(os.path.expanduser("~"), ".docvm_app", "pdf_cache")
        os.makedirs(cache_base, exist_ok=True)
        _pdf_cache_dir = cache_base
    return _pdf_cache_dir


def get_docx_pdf_cache_path(docx_path: str) -> str | None:
    try:
        mtime = os.path.getmtime(docx_path)
        key = f"{docx_path}_{mtime}"
        h = hashlib.md5(key.encode("utf-8")).hexdigest()
        return os.path.join(get_pdf_cache_dir(), f"{h}.pdf")
    except OSError:
        return None


def convert_docx_to_pdf(docx_path: str) -> str | None:
    """用 Word COM 将 .docx 转为 PDF；失败返回 None。带缓存。"""
    if not HAS_WIN32 or not IS_WINDOWS:
        return None

    cache_pdf = get_docx_pdf_cache_path(docx_path)
    if cache_pdf and os.path.exists(cache_pdf):
        return cache_pdf

    word = None
    try:
        word = win32com.client.Dispatch("Word.Application")
        word.Visible = False
        word.DisplayAlerts = False
        abs_docx = os.path.abspath(docx_path)
        doc = word.Documents.Open(abs_docx)
        doc.SaveAs(cache_pdf, FileFormat=17)
        doc.Close(False)
        word.Quit()
        if os.path.exists(cache_pdf):
            return cache_pdf
        return None
    except Exception:
        try:
            if word:
                word.Quit()
        except Exception:
            pass
        return None


def has_word_installed() -> bool:
    global _word_available
    if _word_available is not None:
        return _word_available
    if not HAS_WIN32 or not IS_WINDOWS:
        _word_available = False
        return False
    try:
        word = win32com.client.Dispatch("Word.Application")
        word.Quit()
        _word_available = True
    except Exception:
        _word_available = False
    return _word_available
